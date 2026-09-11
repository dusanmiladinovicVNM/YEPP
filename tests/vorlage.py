#!/usr/bin/env python3
"""
Prueft die Excel-Vorlage gegen echte Ausgabe des Endpunkts.

    python3 tests/vorlage.py

tests/csv_beispiel.mjs erzeugt eine CSV, wie der Endpunkt sie liefert. Die
Formeln werden aus der gebauten Vorlage gelesen, gegen diese Daten ausgewertet
und mit den Quellwerten verglichen. Damit ist geprueft, was hier ueberhaupt
schiefgehen kann: ob die Spaltenbuchstaben und der Suchschluessel stimmen.

Kein Excel im Spiel. Die LibreOffice-Neuberechnung waere der bessere Weg,
aber diese Umgebung hat kein Calc-Modul (libsclo.so und calc.xcd fehlen;
soffice bricht auch bei einer Datei mit drei Zellen mit «source file could
not be loaded» ab). Der Auswerter unten deckt die Formellogik ab — dass
Excel die Datei anstandslos oeffnet, ist damit NICHT gezeigt.
"""

import csv
import io
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WURZEL  = Path(__file__).resolve().parent.parent
VORLAGE = WURZEL / 'vorlage' / 'Wareneingang-Vorlage.xlsx'

bestanden = fehler = 0

# =IFERROR(INDEX(Daten!$F$2:$F$20001,MATCH(<schluessel>,Daten!$A$2:$A$20001,0)),"")
MUSTER = re.compile(
    r'^=IFERROR\(INDEX\(Daten!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+),'
    r'MATCH\((.+?),Daten!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+),0\)\),""\)$'
)

# =IF($J$2="","",IF($J$10=0,"...",IF($J$10><n>,"..."&$J$10&"...","")))
WARNUNG = re.compile(
    r'^=IF\(\$J\$2="","",'
    r'IF\(\$J\$10=0,"([^"]+)",'
    r'IF\(\$J\$10>(\d+),"([^"]*)"&\$J\$10&"([^"]*)",""\)\)\)$'
)


def ok(name, bedingung, extra=''):
    global bestanden, fehler
    if bedingung:
        bestanden += 1
        print(f'  PASS  {name}')
    else:
        fehler += 1
        print(f'  FAIL  {name}' + (f'\n        {extra}' if extra else ''))


def beispiel_csv():
    r = subprocess.run(['node', str(WURZEL / 'tests' / 'csv_beispiel.mjs')],
                       capture_output=True, text=True, cwd=WURZEL)
    if r.returncode:
        sys.exit('csv_beispiel.mjs fehlgeschlagen:\n' + r.stderr)
    nummern = r.stderr.strip().split('\t')[1:]
    return list(csv.reader(io.StringIO(r.stdout))), nummern


class Blatt:
    """Das Blatt `Daten`, wie Power Query es fuellen wuerde."""

    def __init__(self, kopf, zeilen):
        self.kopf = kopf
        # Zeile 1 ist die Ueberschrift, Daten ab Zeile 2 — wie im echten Blatt
        self.zellen = {}
        for i, zeile in enumerate(zeilen, start=2):
            for j, wert in enumerate(zeile, start=1):
                self.zellen[(get_column_letter(j), i)] = wert

    def wert(self, sp, zeile):
        return self.zellen.get((sp, zeile), '')


def auswerten(formel, blatt, wahl, pos_nr, zeilen_da):
    """Rechnet eine Formel der Vorlage aus. Gibt '' zurueck wie IFERROR."""
    m = MUSTER.match(formel)
    if not m:
        raise ValueError(f'unerwartete Formel: {formel}')
    idx_sp, idx_von, idx_sp2, idx_bis, schluessel, m_sp, m_von, m_sp2, m_bis = \
        m.group(1), int(m.group(2)), m.group(3), int(m.group(4)), \
        m.group(5), m.group(6), int(m.group(7)), m.group(8), int(m.group(9))

    # INDEX- und MATCH-Bereich muessen deckungsgleich sein, sonst zeigt der
    # gefundene Zeilenindex auf eine andere Zeile.
    if (idx_von, idx_bis) != (m_von, m_bis):
        raise ValueError(f'Bereiche versetzt: {formel}')
    if idx_sp != idx_sp2 or m_sp != m_sp2:
        raise ValueError(f'Bereich ueber mehrere Spalten: {formel}')

    if schluessel == '$J$2':
        gesucht = wahl
    else:
        s = re.match(r'^\$J\$2&"-"&\$A(\d+)$', schluessel)
        if not s:
            raise ValueError(f'unbekannter Schluessel: {schluessel}')
        gesucht = f'{wahl}-{pos_nr}'

    for zeile in range(idx_von, min(idx_bis, idx_von + zeilen_da) + 1):
        if str(blatt.wert(m_sp, zeile)) == str(gesucht):
            return blatt.wert(idx_sp, zeile)
    return ''                                   # IFERROR faengt #NV ab


def main():
    if not VORLAGE.exists():
        sys.exit('Vorlage fehlt — zuerst tools/vorlage_bauen.py laufen lassen')

    reihen, (nr_voll, nr_offen, nr_storno, nr_viele) = beispiel_csv()
    kopf, daten = reihen[0], reihen[1:]
    blatt = Blatt(kopf, daten)
    n = len(daten)

    wb = load_workbook(VORLAGE)
    fm = wb['Formular']

    print('\n1) Beispieldaten')
    ok('CSV hat 23 Spalten', len(kopf) == 23, str(len(kopf)))
    ok('stornierter Wareneingang fehlt', all(z[0] != nr_storno for z in daten))
    ok('sechzehn Positionszeilen', n == 16, str(n))

    print('\n2) Aufbau der Arbeitsmappe')
    ok('vier Blaetter', wb.sheetnames == ['Formular', 'Daten', 'Nummern', 'Liste'],
       str(wb.sheetnames))
    ok('Daten traegt die CSV-Ueberschriften',
       [wb['Daten'].cell(row=1, column=i + 1).value for i in range(23)] == kopf)
    ok('Liste traegt dieselben Ueberschriften',
       [wb['Liste'].cell(row=1, column=i + 1).value for i in range(23)] == kopf)
    ok('Auswahlliste haengt an J2',
       any('J2' in str(dv.sqref) and dv.formula1 == 'Nummern!$A$2:$A$1000'
           for dv in fm.data_validations.dataValidation),
       str([(str(d.sqref), d.formula1) for d in fm.data_validations.dataValidation]))
    # openpyxl setzt den Blattnamen in Anfuehrungszeichen
    druck = str(fm.print_area).replace("'", '').strip('[]')
    ok('Druckbereich endet bei Spalte H', druck == 'Formular!$A$1:$H$30', druck)
    # J2 selbst liegt ausserhalb des Druckbereichs; ohne F1 traegt das
    # unterschriebene Blatt keine Nummer.
    ok('Nummer im Druckbereich', fm['F1'].value == '=IF($J$2="","",$J$2)',
       repr(fm['F1'].value))
    ok('Nummer ueber F1:H1 verbunden',
       'F1:H1' in {str(b) for b in fm.merged_cells.ranges})
    ok('auf eine Seite skaliert', fm.sheet_properties.pageSetUpPr.fitToPage is True)
    ok('Bestehend-Spalte gelb',
       all(fm[f'H{z}'].fill.fgColor.rgb.endswith('FFFF00') for z in range(16, 26)))

    print('\n3) Papierraster')
    raster = {
        'A1':  'Wareneingang / Material reception',
        'A3':  'Aufgabe / Task',
        'C3':  'Name Mitarbeiter / Employee name',
        'E3':  'Datum / Date',
        'F3':  'Uhrzeit / Time',
        'A6':  'Eingelagert / stored',
        'A10': 'Kunde / Client',
        'A11': 'Lieferant / Supplier',
        'A15': 'N°',
        'C15': 'Anzahl / QTY',
        'D15': 'kg',
        'E15': 'MHD / Expiration Date',
        'H15': 'Bestehend / Existing',
        'A27': 'Lagerfläche / storage space',
        'A28': 'm2 Anzahl / m2 quantity',
        'D28': '1 g = 0.001 kg',
    }
    for zelle, text in raster.items():
        ok(f'{zelle} = {text[:34]}', fm[zelle].value == text, repr(fm[zelle].value))
    ok('Positionen 1..10 durchnummeriert',
       [fm[f'A{z}'].value for z in range(16, 26)] == list(range(1, 11)))

    # -----------------------------------------------------------------
    print('\n4) Formeln gegen die Quelldaten — vollstaendig quittiert')
    quelle = {k: v for k, v in zip(kopf, daten[0])}
    W = lambda z, nr=None: auswerten(fm[z].value, blatt, nr_voll, nr, n)  # noqa: E731

    pruefungen = [
        ('C10', None, 'Kunde'), ('C11', None, 'Lieferant'),
        ('C4', None, 'AngNam'), ('E4', None, 'AngDat'), ('F4', None, 'AngZeit'),
        ('C5', None, 'GezNam'), ('E5', None, 'GezDat'), ('F5', None, 'GezZeit'),
        ('C6', None, 'EinNam'), ('E6', None, 'EinDat'), ('F6', None, 'EinZeit'),
        ('C28', None, 'LagerM2'), ('B30', None, 'KopfBemerkung'),
    ]
    for zelle, nr, feld in pruefungen:
        ok(f'{zelle} zieht {feld}', W(zelle, nr) == quelle[feld],
           f'{W(zelle, nr)!r} statt {quelle[feld]!r}')

    ok('Kunde mit Komma unzerteilt', W('C10') == 'Meier, Sohn & Co', repr(W('C10')))
    ok('Gezaehlt zeigt den anderen Mitarbeiter', W('C5') == 'Bob Meier', repr(W('C5')))

    # Ohne die Verbindungen stehen «Gezählt & kontrolliert / counted &
    # controlled» und der Name in der 9 Zeichen schmalen N°-Spalte.
    verbunden = {str(b) for b in fm.merged_cells.ranges}
    for bereich in ('A3:B3', 'C3:D3', 'A5:B5', 'C5:D5',
                    'A10:B10', 'C10:F10', 'A11:B11', 'C11:F11',
                    'A8:H8', 'A13:H13'):
        ok(f'{bereich} verbunden', bereich in verbunden, str(sorted(verbunden)))

    print('\n5) Positionen')
    pos = [dict(zip(kopf, z)) for z in daten if z[0] == nr_voll]
    ok('drei Positionen', len(pos) == 3, str(len(pos)))
    felder = [('B', 'Artikel'), ('C', 'Anzahl'), ('D', 'KG'), ('E', 'MHD'),
              ('F', 'Regalplatz'), ('G', 'Bemerkung'), ('H', 'Bestehend')]
    for i, p in enumerate(pos):
        zeile = 16 + i
        for sp, feld in felder:
            ok(f'{sp}{zeile} = {feld}',
               auswerten(fm[f'{sp}{zeile}'].value, blatt, nr_voll, i + 1, n) == p[feld],
               f'{auswerten(fm[f"{sp}{zeile}"].value, blatt, nr_voll, i + 1, n)!r} '
               f'statt {p[feld]!r}')

    ok('Anfuehrungszeichen im Artikel',
       auswerten(fm['B17'].value, blatt, nr_voll, 2, n) == 'Rohr "40mm"')
    ok('Bestehend gesetzt',
       auswerten(fm['H16'].value, blatt, nr_voll, 1, n) == 'X')
    ok('Bestehend nicht gesetzt',
       auswerten(fm['H17'].value, blatt, nr_voll, 2, n) == '')

    print('\n6) Ungenutzte Zeilen bleiben leer')
    for i, zeile in enumerate(range(19, 26), start=4):
        leer = [sp for sp, _ in felder
                if auswerten(fm[f'{sp}{zeile}'].value, blatt, nr_voll, i, n) != '']
        ok(f'Zeile {zeile} leer', not leer, str(leer))

    print('\n7) Anderer Wareneingang')
    O = lambda z, nr=None: auswerten(fm[z].value, blatt, nr_offen, nr, n)  # noqa: E731
    ok('Kunde gewechselt', O('C10') == 'Zweiter Kunde AG', repr(O('C10')))
    ok('Lieferant gewechselt', O('C11') == 'Alpina Food', repr(O('C11')))
    ok('erste Position gewechselt',
       auswerten(fm['B16'].value, blatt, nr_offen, 1, n) == 'Mehl Type 550')
    ok('nicht quittierter Schritt leer', O('C5') == '', repr(O('C5')))
    ok('leeres LagerM2 leer', O('C28') == '', repr(O('C28')))
    ok('zweite Position leer',
       auswerten(fm['B17'].value, blatt, nr_offen, 2, n) == '')

    print('\n8) Unbekannte und stornierte Nummer')
    for nummer, wie in ((nr_storno, 'storniert'), ('WE-1999-9999', 'unbekannt')):
        belegt = [z for z in ('C4', 'C5', 'C6', 'C10', 'C11', 'C28', 'B30')
                  if auswerten(fm[z].value, blatt, nummer, None, n) != '']
        belegt += [f'B{16 + i}' for i in range(3)
                   if auswerten(fm[f'B{16 + i}'].value, blatt, nummer, i + 1, n) != '']
        ok(f'{wie}: alles leer, kein #NV', not belegt, str(belegt))

    print('\n9) Warnzeile — das Formularblatt hat feste Zeilen')
    # Der Zaehler muss denselben Bereich absuchen wie die uebrigen Formeln,
    # sonst warnt er ueber andere Daten, als das Blatt anzeigt.
    m_kopf = MUSTER.match(fm['C10'].value)
    we_bereich = (f'Daten!${m_kopf.group(6)}${m_kopf.group(7)}:'
                  f'${m_kopf.group(8)}${m_kopf.group(9)}')
    ok('Zaehler nutzt den WeNr-Bereich der Formeln',
       fm['J10'].value == f'=COUNTIF({we_bereich},$J$2)', repr(fm['J10'].value))
    ok('Zaehler steht ausserhalb des Druckbereichs', druck.endswith('$H$30'))

    w = WARNUNG.match(str(fm['A26'].value))
    ok('Warnzeile hat die erwartete Form', w is not None, repr(fm['A26'].value))
    grenze = int(w.group(2))
    zeilen_im_blatt = sum(1 for z in range(16, 40) if fm[f'A{z}'].value in range(1, 99))
    ok('Grenze ist die Zahl der Formularzeilen', grenze == zeilen_im_blatt,
       f'{grenze} statt {zeilen_im_blatt}')

    def warnung(nummer):
        anzahl = sum(1 for z in daten if z[0] == nummer)
        if anzahl == 0:
            return w.group(1)
        if anzahl > grenze:
            return w.group(3) + str(anzahl) + w.group(4)
        return ''

    ok('kurzer Wareneingang ohne Warnung', warnung(nr_voll) == '', warnung(nr_voll))
    ok('zwoelf Positionen im Beispiel',
       sum(1 for z in daten if z[0] == nr_viele) == 12)
    ok('langer Wareneingang warnt mit Zahl',
       '12' in warnung(nr_viele) and str(grenze) in warnung(nr_viele),
       warnung(nr_viele))
    ok('unbekannte Nummer meldet fehlende Daten',
       warnung('WE-1999-9999') == w.group(1), warnung('WE-1999-9999'))

    # Ohne die Warnzeile bliebe genau das hier unbemerkt:
    sichtbar = [auswerten(fm[f'B{16 + i}'].value, blatt, nr_viele, i + 1, n)
                for i in range(10)]
    ok('Blatt zeigt die ersten zehn Positionen',
       sichtbar == [f'Palette {i + 1}' for i in range(10)], str(sichtbar[:3]))
    ok('elfte Position erscheint nirgends',
       'Palette 11' not in [fm[f'B{z}'].value for z in range(16, 26)])

    # ---------------------------------------------------------------
    # Der M-Code der beiden Abfragen
    # ---------------------------------------------------------------
    # Auf dem Mac legt kein Makro die Abfragen an; dort wird dieser Text
    # eingefuegt. Er muss deshalb allein stehen koennen — und mit dem im
    # Makro uebereinstimmen, sonst laden Windows und Mac verschieden.
    m_text = (WURZEL / 'vorlage' / 'Abfragen.m').read_text(encoding='utf-8')
    bas    = (WURZEL / 'vorlage' / 'Vorlage-Aufbau.bas').read_text(encoding='utf-8')

    # Nur der Daten-Teil: die Nummernabfrage enthaelt {"WeNr", Order.Descending},
    # und das ist eine Sortierung, keine Typangabe.
    m_daten = m_text.split('// ============ Nummern')[0]
    getypt = dict(re.findall(r'\{"([^"]+)", ([^}]+)\}', m_daten))
    ok('jede CSV-Spalte ist ausdruecklich getypt',
       [n for n in kopf if n not in getypt] == [],
       str([n for n in kopf if n not in getypt]))
    ok('und keine erfundene dazu',
       [n for n in getypt if n not in kopf] == [],
       str([n for n in getypt if n not in kopf]))

    zahlen = {'Anzahl': 'type number', 'KG': 'type number',
              'LagerM2': 'type number', 'Nr': 'Int64.Type'}
    for name, typ in zahlen.items():
        ok(f'{name} ist eine Zahl', getypt.get(name) == typ, str(getypt.get(name)))

    # Die Datums- und Uhrzeitspalten sind der eigentliche Grund fuer die
    # Liste: als Datum geladen verschieben sie sich um die Zeitzone, und im
    # Formular stuende dann ein Tag daneben.
    datums = [n for n in kopf if n.endswith('Dat') or n.endswith('Zeit')
              or n == 'MHD']
    ok('Datum und Uhrzeit bleiben Text',
       all(getypt.get(n) == 'type text' for n in datums),
       str({n: getypt.get(n) for n in datums if getypt.get(n) != 'type text'}))
    ok('es sind ueberhaupt welche darunter', len(datums) >= 7, str(len(datums)))

    ok('das Gebietsschema steht dabei', '"en-US"' in m_text)
    ok('die Nummernabfrage baut auf Daten auf',
       'Quelle = Daten' in m_text and 'Table.Distinct' in m_text)
    ok('und sortiert absteigend', 'Order.Descending' in m_text)
    # Auf dem Mac laedt eine Abfrage in ein Ziel; «Liste» braucht darum eine
    # eigene, die auf «Daten» zeigt. Windows haengt dieselbe Abfrage zweimal
    # ein — in der Arbeitsmappe steht danach dasselbe.
    for name in ('Daten', 'Nummern', 'Liste'):
        ok(f'der Block «{name}» steht in der Datei',
           f'============ {name} ============' in m_text)

    # Windows (Makro) und Mac (eingefuegt) muessen denselben Text laden.
    # VBA: 1024 Zeichen und 24 Fortsetzungen je LOGISCHER Zeile. Die alte
    # Fassung stand bei 1003 — eine Spalte mehr haette den Import unter
    # Windows zerbrochen, waehrend der Mac-Weg weitergelaufen waere.
    logisch, fort, laengste, meiste = '', 0, 0, 0
    for roh in bas.replace('\r\n', '\n').split('\n'):
        if roh.rstrip().endswith(' _'):
            logisch += roh.rstrip()[:-1]
            fort += 1
            continue
        logisch += roh
        laengste = max(laengste, len(logisch))
        meiste = max(meiste, fort)
        logisch, fort = '', 0
    ok('keine VBA-Zeile nahe an der 1024er-Grenze', laengste < 1000, f'{laengste} Zeichen')
    ok('und keine mit zu vielen Fortsetzungen', meiste <= 24, f'{meiste}')

    bas_typen = dict(re.findall(r'\{""([^"]+)"", ((?:type \w+|Int64\.Type))\}', bas))
    ok('Makro und M-Datei typen gleich', bas_typen == getypt,
       str({n: (bas_typen.get(n), getypt.get(n))
            for n in set(list(bas_typen) + list(getypt))
            if bas_typen.get(n) != getypt.get(n)}))

    print('\n' + '=' * 46)
    print(f'{bestanden} bestanden, {fehler} gescheitert')
    return 1 if fehler else 0


if __name__ == '__main__':
    sys.exit(main())
