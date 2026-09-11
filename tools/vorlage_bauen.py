#!/usr/bin/env python3
"""
Baut die Excel-Vorlage `vorlage/Wareneingang-Vorlage.xlsx`.

    python3 tools/vorlage_bauen.py

Die Spaltenreihenfolge wird aus `CSV_SPALTEN` in apps-script/Code.gs gelesen,
nicht hier wiederholt. Aendert sich der Endpunkt, wird die Vorlage neu gebaut
und kann nicht auseinanderlaufen — genau der Fehler, den eine von Hand
gepflegte Vorlagendatei macht.

Was NICHT hier entsteht: die Power-Query-Teile IN der Datei. Die legt Excel
selbst an — von aussen erzeugte lehnt es gern wortlos ab, und eine Datei,
die ohne Meldung nichts tut, ist schlimmer als ein Handgriff.

Was hier sehr wohl entsteht, ist der M-Code der beiden Abfragen, in
`vorlage/Abfragen.m` und zugleich im Makro `Vorlage-Aufbau.bas`. Windows
laesst ihn das Makro anlegen, der Mac fuegt ihn im erweiterten Editor ein —
aber es ist derselbe Text aus derselben Quelle, und damit koennen die
beiden Wege nicht auseinanderlaufen.
"""

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

WURZEL  = Path(__file__).resolve().parent.parent
CODE_GS = WURZEL / 'apps-script' / 'Code.gs'
ZIEL    = WURZEL / 'vorlage' / 'Wareneingang-Vorlage.xlsx'

SCHRIFT   = 'Arial'
POS_ZEILEN = 10          # Formularzeilen; das Papier hat 5, 10 deckt mehr ab
KOPF_ZEILE = 15          # Zeile der Positionsueberschrift, wie in blattAufbauen
DATEN_ZEILEN = 20000     # Suchbereich in `Daten`

# Begrenzte Bereiche statt ganzer Spalten (Daten!A:A). Excel optimiert ganze
# Spalten auf den benutzten Bereich, andere Rechner nicht — dort laeuft jede
# der rund achtzig Formeln ueber eine Million Zeilen. Reicht der Bereich
# nicht mehr, hier erhoehen und neu bauen; `&tage=` am Endpunkt haelt die
# Datenmenge ohnehin klein.

ROT   = 'C00000'         # Warnzeile, nur sichtbar wenn sie etwas zu sagen hat

GRAU  = PatternFill('solid', fgColor='F2F2F2')
GELB  = PatternFill('solid', fgColor='FFFF00')
BLASS = PatternFill('solid', fgColor='FFF9E3')      # Eingabefeld
DUENN = Side(style='thin', color='000000')
RAHMEN = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)


def csv_spalten():
    """Liest CSV_SPALTEN aus Code.gs — eine Quelle, nicht zwei."""
    text = CODE_GS.read_text(encoding='utf-8')
    m = re.search(r'const CSV_SPALTEN = \[(.*?)\];', text, re.S)
    if not m:
        sys.exit('CSV_SPALTEN in Code.gs nicht gefunden')
    namen = re.findall(r"'([^']+)'", m.group(1))
    if len(namen) < 23:
        sys.exit(f'CSV_SPALTEN wirkt unvollstaendig: {namen}')
    return namen


def spalte(namen, name):
    """Spaltenbuchstabe im Blatt `Daten` fuer einen CSV-Feldnamen."""
    return get_column_letter(namen.index(name) + 1)


def bereich(namen, name):
    """Begrenzter Suchbereich einer Spalte, ab Zeile 2."""
    sp = spalte(namen, name)
    return f'Daten!${sp}$2:${sp}${DATEN_ZEILEN + 1}'


def setz(bl, zelle, wert, fett=False, gross=None, wrap=False,
         rahmen=False, fuell=None, aus=None, oben='center'):
    z = bl[zelle]
    z.value = wert
    z.font = Font(name=SCHRIFT, bold=fett, size=gross or 10)
    z.alignment = Alignment(wrap_text=wrap, vertical=oben, horizontal=aus)
    if rahmen:
        z.border = RAHMEN
    if fuell:
        z.fill = fuell
    return z


def rahmen_um(bl, bereich):
    for reihe in bl[bereich]:
        for z in reihe:
            z.border = RAHMEN


# Die vier Spalten, die wirklich Zahlen sind. Alles andere bleibt Text —
# besonders die Datums- und Uhrzeitspalten: der Endpunkt liefert sie als
# `GGGG-MM-TT` und `HH:MM`, und als Datum geladen verschieben sie sich um
# die Zeitzone.
ZAHLEN = {'Anzahl': 'type number', 'KG': 'type number',
          'LagerM2': 'type number', 'Nr': 'Int64.Type'}

M_ZIEL = WURZEL / 'vorlage' / 'Abfragen.m'
BAS    = WURZEL / 'vorlage' / 'Vorlage-Aufbau.bas'


def m_typen(namen):
    """
    JEDE Spalte wird ausdruecklich getypt, nicht nur die vier Zahlen.

    Ohne das haengt es an Excel, was aus `AngDat` wird, und die Antwort
    faellt je nach Fassung anders aus. Ausgeschrieben steht es fest — und
    weil die Liste aus CSV_SPALTEN kommt, kann eine neue Spalte nicht
    stillschweigend ungetypt bleiben.
    """
    teile = [f'{{"{n}", {ZAHLEN.get(n, "type text")}}}' for n in namen]
    return ', '.join(teile)


def m_abfragen(namen, quelle='<Web-App-URL>?token=<TOKEN_READ>&format=csv&tage=365'):
    daten = (
        'let\n'
        f'    Quelle = Csv.Document(Web.Contents("{quelle}"),'
        '[Delimiter=",", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),\n'
        '    Kopf = Table.PromoteHeaders(Quelle, [PromoteAllScalars=true]),\n'
        f'    Typen = Table.TransformColumnTypes(Kopf,{{{m_typen(namen)}}}, "en-US")\n'
        'in\n'
        '    Typen'
    )
    nummern = (
        'let\n'
        '    Quelle = Daten,\n'
        '    NurNr = Table.SelectColumns(Quelle,{"WeNr"}),\n'
        '    Eindeutig = Table.Distinct(NurNr),\n'
        '    Sortiert = Table.Sort(Eindeutig,{{"WeNr", Order.Descending}})\n'
        'in\n'
        '    Sortiert'
    )
    # Auf dem Mac laedt eine Abfrage in EIN Ziel. Das Makro haengt unter
    # Windows dieselbe Abfrage zweimal ein (Daten und Liste); durch die
    # Oberflaeche geht das nicht, also bekommt `Liste` eine eigene Abfrage,
    # die nichts weiter tut als auf `Daten` zu zeigen. In der Arbeitsmappe
    # steht danach dasselbe — nur der Weg dorthin unterscheidet sich, weil
    # die beiden Excel-Fassungen sich unterscheiden.
    liste = 'let\n    Quelle = Daten\nin\n    Quelle'
    return daten, nummern, liste


def m_schreiben(namen):
    """Der M-Code zum Einfuegen — und derselbe Text ins Makro."""
    daten, nummern, liste = m_abfragen(namen)
    M_ZIEL.write_text(
        '// Power-Query-Abfragen fuer Wareneingang-Vorlage.xlsx\n'
        '// ERZEUGT von tools/vorlage_bauen.py — nicht von Hand aendern.\n'
        '//\n'
        '// Mac: Daten -> Daten abrufen -> Leere Abfrage, dann\n'
        '// Erweiterter Editor, alles ersetzen, Abfrage genau so nennen wie\n'
        '// die Ueberschrift hier. Dreimal, in dieser Reihenfolge — Nummern\n'
        '// und Liste bauen auf Daten auf.\n'
        '//\n'
        '// Zwei Platzhalter, und beide NUR im Block «Daten»:\n'
        '//   <Web-App-URL>  die Adresse der Bereitstellung, die auf /exec\n'
        '//                  endet — dieselbe, mit der die App spricht.\n'
        '//   <TOKEN_READ>   schreibt einrichtungPruefen() im Apps Script\n'
        '//                  fertig hin, in der letzten Zeile des Berichts.\n'
        '// Nummern und Liste bekommen keine Adresse: sie lesen Daten.\n'
        '//\n'
        '// ============ Daten ============\n'
        + daten +
        '\n\n// ============ Nummern ============\n'
        + nummern +
        '\n\n// ============ Liste ============\n'
        + liste + '\n', encoding='utf-8')

    # Die GANZE Funktion neu schreiben, nicht ein Stueck darin ersetzen.
    #
    # VBA laesst je LOGISCHER Zeile — Fortsetzungen zusammengerechnet —
    # 1024 Zeichen zu. Mit einem Typ je CSV-Spalte stand die alte Fassung
    # bei 1003: einundzwanzig Zeichen vor der Grenze. Die naechste Spalte
    # haette den Import unter Windows zerbrochen, waehrend der Mac-Weg
    # (Abfragen.m, ohne diese Grenze) weitergelaufen waere — genau das
    # Auseinanderlaufen, das hier nicht passieren soll.
    #
    # Mehr Fortsetzungen helfen dagegen nicht: sie gehoeren zu derselben
    # logischen Zeile. Also mehrere ANWEISUNGEN, jede fuer sich kurz.
    zitat = lambda t: '"' + t.replace('"', '""') + '"'

    teile = [f'{{"{n}", {ZAHLEN.get(n, "type text")}}}' for n in namen]
    stuecke, zeile = [], ''
    for t in teile:
        kandidat = (zeile + ', ' + t) if zeile else t
        if len(kandidat) > 120:
            stuecke.append(zeile + ', ')
            zeile = t
        else:
            zeile = kandidat
    stuecke.append(zeile)

    zeilen = [
        'Private Function MDaten(ByVal quelle As String) As String',
        "    ' ERZEUGT von tools/vorlage_bauen.py — nicht von Hand aendern.",
        "    ' Mehrere Anweisungen statt einer langen: VBA laesst je logischer",
        "    ' Zeile nur 1024 Zeichen zu, und die Typenliste waechst mit den",
        "    ' Spalten. Derselbe Text steht in vorlage/Abfragen.m.",
        '    Dim m As String',
        '    m = "let" & vbLf',
        '    m = m & "    Quelle = Csv.Document(Web.Contents(""" & quelle & """),"',
        '    m = m & "[Delimiter="","", Encoding=65001, '
        'QuoteStyle=QuoteStyle.Csv])," & vbLf',
        '    m = m & "    Kopf = Table.PromoteHeaders(Quelle, '
        '[PromoteAllScalars=true])," & vbLf',
        '    m = m & "    Typen = Table.TransformColumnTypes(Kopf,{"',
    ]
    zeilen += [f'    m = m & {zitat(st)}' for st in stuecke]
    zeilen += [
        '    m = m & "}, ""en-US"")" & vbLf',
        '    m = m & "in" & vbLf & "    Typen"',
        '    MDaten = m',
        'End Function',
    ]

    text = BAS.read_text(encoding='utf-8')
    neu = re.sub(r'Private Function MDaten\(.*?End Function',
                 lambda _: '\r\n'.join(zeilen), text, flags=re.S)
    if neu == text:
        sys.exit('MDaten liess sich nicht ersetzen — '
                 'Vorlage-Aufbau.bas hat sich geaendert.')
    BAS.write_text(neu, encoding='utf-8')
    vba_grenzen_pruefen(neu)


def vba_grenzen_pruefen(text):
    """
    VBA: hoechstens 1024 Zeichen und 24 Fortsetzungen je logischer Zeile.

    Ohne diese Pruefung faellt das Ueberschreiten erst beim Importieren in
    den VBA-Editor auf — auf einem Rechner, der hier gar nicht steht.
    """
    logisch, fort = '', 0
    for roh in text.replace('\r\n', '\n').split('\n'):
        if roh.rstrip().endswith(' _'):
            logisch += roh.rstrip()[:-1]
            fort += 1
            continue
        logisch += roh
        if len(logisch) > 1000:
            sys.exit(f'VBA-Zeile zu lang ({len(logisch)} Zeichen, Grenze 1024): '
                     f'{logisch[:80]}…')
        if fort > 24:
            sys.exit(f'Zu viele Fortsetzungen ({fort}, Grenze 24)')
        logisch, fort = '', 0


def bauen():
    namen = csv_spalten()
    B = lambda n: bereich(namen, n)         # noqa: E731
    WE  = B('WeNr')
    SCH = B('Schluessel')

    wb = Workbook()

    # ---------------------------------------------------------------
    # Formular — das Papier, gefuellt aus `Daten`
    # ---------------------------------------------------------------
    fm = wb.active
    fm.title = 'Formular'

    for sp, breite in zip('ABCDEFGHIJ', [9, 34, 12, 11, 17, 20, 22, 13, 3, 24]):
        fm.column_dimensions[sp].width = breite

    setz(fm, 'A1', 'Wareneingang / Material reception', fett=True, gross=14)

    # Die Nummer gehoert in den Druckbereich: der Ausdruck wird
    # unterschrieben und abgelegt. J2 selbst steht ausserhalb.
    setz(fm, 'F1', '=IF($J$2="","",$J$2)', fett=True, gross=12, aus='right')
    fm.merge_cells('F1:H1')

    # Auswahl — ausserhalb des Druckbereichs, damit sie nicht mitgedruckt wird
    setz(fm, 'J1', 'Wareneingang wählen', fett=True)
    wahl = setz(fm, 'J2', '', fett=True, rahmen=True, fuell=BLASS)
    wahl.font = Font(name=SCHRIFT, bold=True, size=11)
    setz(fm, 'J4', 'Nur diese Zelle wird von Hand '
                   'geändert. Alles andere füllt sich selbst.', wrap=True)
    fm['J4'].alignment = Alignment(wrap_text=True, vertical='top')
    fm.merge_cells('J4:J7')

    setz(fm, 'J9', 'Positionen im Dokument', gross=8)
    setz(fm, 'J10', f'=COUNTIF({B("WeNr")},$J$2)', fett=True)

    pruef = DataValidation(type='list', formula1='Nummern!$A$2:$A$1000',
                           allow_blank=True, showDropDown=False)
    pruef.error = 'Diese Nummer steht nicht in der Liste.'
    pruef.prompt = 'Wareneingang aus der Liste wählen'
    fm.add_data_validation(pruef)
    pruef.add(fm['J2'])

    # -- Quittungen -------------------------------------------------
    # Spalte A ist die schmale N°-Spalte des Positionsblocks; die
    # Beschriftungen des Papiers passen dort nicht hinein und wuerden bei
    # fester Zeilenhoehe abgeschnitten. Darum vier verbundene Bereiche:
    # A:B Aufgabe, C:D Name, E Datum, F Uhrzeit — wie in blattAufbauen.
    for sp, titel in zip('ACEF', ['Aufgabe / Task',
                                  'Name Mitarbeiter / Employee name',
                                  'Datum / Date', 'Uhrzeit / Time']):
        setz(fm, f'{sp}3', titel, fett=True, wrap=True, rahmen=True, fuell=GRAU)
    fm.row_dimensions[3].height = 26

    schritte = [
        (4, 'Angenommen / Accepted',                         'AngNam', 'AngDat', 'AngZeit'),
        (5, 'Gezählt & kontrolliert / counted & controlled', 'GezNam', 'GezDat', 'GezZeit'),
        (6, 'Eingelagert / stored',                          'EinNam', 'EinDat', 'EinZeit'),
    ]
    for zeile, beschriftung, nam, dat, zeit in schritte:
        setz(fm, f'A{zeile}', beschriftung, fett=True, wrap=True, rahmen=True)
        for sp, feld in zip('CEF', (nam, dat, zeit)):
            setz(fm, f'{sp}{zeile}',
                 f'=IFERROR(INDEX({B(feld)},MATCH($J$2,{WE},0)),"")',
                 rahmen=True)
        fm.row_dimensions[zeile].height = 26

    for zeile in range(3, 7):
        fm.merge_cells(f'A{zeile}:B{zeile}')
        fm.merge_cells(f'C{zeile}:D{zeile}')
        rahmen_um(fm, f'A{zeile}:F{zeile}')

    setz(fm, 'A8',
         'Artikelanzahl bitte direkt auf dem Lieferschein abhaken bzw. anpassen.\n'
         'Please check off or adapt the article quantities directly on the delivery slip',
         fett=True, wrap=True, oben='top')
    fm.merge_cells('A8:H8')
    fm.row_dimensions[8].height = 28

    # -- Kunde / Lieferant ------------------------------------------
    for zeile, beschriftung, feld in [(10, 'Kunde / Client',      'Kunde'),
                                      (11, 'Lieferant / Supplier', 'Lieferant')]:
        setz(fm, f'A{zeile}', beschriftung, fett=True, rahmen=True)
        setz(fm, f'C{zeile}',
             f'=IFERROR(INDEX({B(feld)},MATCH($J$2,{WE},0)),"")', rahmen=True)
        fm.merge_cells(f'A{zeile}:B{zeile}')
        fm.merge_cells(f'C{zeile}:F{zeile}')
        rahmen_um(fm, f'A{zeile}:F{zeile}')

    setz(fm, 'A13',
         'Bei neuem und bestehendem Material mit oder ohne Lieferschein notwendig:\n'
         'For new and existing material with or without delivery slip needed',
         fett=True, wrap=True, oben='top')
    fm.merge_cells('A13:H13')
    fm.row_dimensions[13].height = 28

    # -- Positionen -------------------------------------------------
    ueberschriften = [
        'N°', 'Artikelbezeichnung / Article description', 'Anzahl / QTY', 'kg',
        'MHD / Expiration Date',
        'Regalplatznr., eralten durch Secend Team/shelf no. received by Secend Team',
        'Bemerkungen /Kontrolle\nComments / Checks', 'Bestehend / Existing',
    ]
    for sp, titel in zip('ABCDEFGH', ueberschriften):
        setz(fm, f'{sp}{KOPF_ZEILE}', titel, fett=True, wrap=True,
             rahmen=True, fuell=GRAU, oben='bottom')
    fm.row_dimensions[KOPF_ZEILE].height = 58

    pos_felder = [('B', 'Artikel'), ('C', 'Anzahl'), ('D', 'KG'), ('E', 'MHD'),
                  ('F', 'Regalplatz'), ('G', 'Bemerkung'), ('H', 'Bestehend')]

    for i in range(POS_ZEILEN):
        zeile = KOPF_ZEILE + 1 + i
        setz(fm, f'A{zeile}', i + 1, rahmen=True, fuell=GRAU, aus='center')
        for sp, feld in pos_felder:
            # Schluessel WeNr-Nr: ein gewoehnlicher VERGLEICH statt einer
            # Matrixformel ueber zwei Kriterien — laeuft in jedem Excel.
            setz(fm, f'{sp}{zeile}',
                 f'=IFERROR(INDEX({B(feld)},'
                 f'MATCH($J$2&"-"&$A{zeile},{SCH},0)),"")',
                 rahmen=True, wrap=(sp in 'BG'),
                 aus='right' if sp in 'CD' else ('center' if sp == 'H' else None),
                 fuell=GELB if sp == 'H' else None)
        fm.row_dimensions[zeile].height = 30

    # -- Warnzeile --------------------------------------------------
    # Das Formularblatt hat feste Zeilen; ein laengerer Wareneingang wuerde
    # sonst stillschweigend abgeschnitten und der Ausdruck saehe vollstaendig
    # aus. Dieselbe Zeile meldet auch den umgekehrten Fall: gar keine Daten,
    # weil die Abfrage nicht aktualisiert oder der Suchbereich zu klein ist.
    warn = KOPF_ZEILE + POS_ZEILEN + 1
    setz(fm, f'A{warn}',
         f'=IF($J$2="","",'
         f'IF($J$10=0,"Zu dieser Nummer stehen keine Zeilen im Blatt Daten — '
         f'aktualisieren, oder der Suchbereich reicht nicht.",'
         f'IF($J$10>{POS_ZEILEN},"Achtung: dieser Wareneingang hat "&$J$10&'
         f'" Positionen. Hier stehen nur die ersten {POS_ZEILEN}.","")))',
         fett=True, wrap=True)
    fm[f'A{warn}'].font = Font(name=SCHRIFT, bold=True, size=10, color=ROT)
    fm.merge_cells(f'A{warn}:H{warn}')

    # -- Fuss -------------------------------------------------------
    fuss = KOPF_ZEILE + POS_ZEILEN + 2
    setz(fm, f'A{fuss}', 'Lagerfläche / storage space', fett=True, rahmen=True)
    fm.merge_cells(f'A{fuss}:B{fuss}')
    setz(fm, f'C{fuss}', '', rahmen=True)
    setz(fm, f'D{fuss}', 'Umrechung/Conversion', fett=True, rahmen=True)

    setz(fm, f'A{fuss + 1}', 'm2 Anzahl / m2 quantity', fett=True, rahmen=True)
    fm.merge_cells(f'A{fuss + 1}:B{fuss + 1}')
    setz(fm, f'C{fuss + 1}',
         f'=IFERROR(INDEX({B("LagerM2")},MATCH($J$2,{WE},0)),"")',
         rahmen=True, aus='right')
    setz(fm, f'D{fuss + 1}', '1 g = 0.001 kg', fett=True, rahmen=True)
    rahmen_um(fm, f'A{fuss}:D{fuss + 1}')

    setz(fm, f'A{fuss + 3}', 'Bemerkung', fett=True)
    setz(fm, f'B{fuss + 3}',
         f'=IFERROR(INDEX({B("KopfBemerkung")},MATCH($J$2,{WE},0)),"")')
    fm.merge_cells(f'B{fuss + 3}:H{fuss + 3}')

    fm.print_area = f'A1:H{fuss + 3}'
    fm.page_setup.orientation = 'portrait'
    fm.page_setup.paperSize = fm.PAPERSIZE_A4
    fm.page_setup.fitToWidth = 1
    fm.page_setup.fitToHeight = 1
    fm.sheet_properties.pageSetUpPr.fitToPage = True
    fm.freeze_panes = 'A16'

    # ---------------------------------------------------------------
    # Daten — Ziel der Power-Query-Abfrage
    # ---------------------------------------------------------------
    dt = wb.create_sheet('Daten')
    for i, name in enumerate(namen, start=1):
        setz(dt, f'{get_column_letter(i)}1', name, fett=True, fuell=GRAU)
        dt.column_dimensions[get_column_letter(i)].width = max(11, len(name) + 3)
    dt.freeze_panes = 'A2'

    # ---------------------------------------------------------------
    # Nummern — Quelle der Auswahlliste
    # ---------------------------------------------------------------
    nm = wb.create_sheet('Nummern')
    setz(nm, 'A1', 'WeNr', fett=True, fuell=GRAU)
    nm.column_dimensions['A'].width = 18
    nm.freeze_panes = 'A2'

    # ---------------------------------------------------------------
    # Liste — dieselbe Abfrage, nur sichtbar
    # ---------------------------------------------------------------
    ls = wb.create_sheet('Liste')
    for i, name in enumerate(namen, start=1):
        setz(ls, f'{get_column_letter(i)}1', name, fett=True, fuell=GRAU)
        ls.column_dimensions[get_column_letter(i)].width = max(11, len(name) + 3)
    ls.freeze_panes = 'A2'
    ls.auto_filter.ref = f'A1:{get_column_letter(len(namen))}1'

    ZIEL.parent.mkdir(exist_ok=True)
    wb.save(ZIEL)
    m_schreiben(namen)
    print(f'{ZIEL.relative_to(WURZEL)} gebaut — {len(namen)} Spalten, '
          f'{POS_ZEILEN} Positionszeilen, Warnzeile in A{warn}')
    print(f'{M_ZIEL.relative_to(WURZEL)} geschrieben — '
          f'{len(namen)} Spalten ausdruecklich getypt')
    return namen


if __name__ == '__main__':
    bauen()
